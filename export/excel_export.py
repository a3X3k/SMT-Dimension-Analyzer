from openpyxl import Workbook
from openpyxl.styles import Font
from datetime import datetime

def export_excel(path, unique_parts, cad_by_ref=None, dimension_results=None, shape_models=None):
    cad_by_ref=cad_by_ref or {}; dimension_results=dimension_results or {}; shape_models=shape_models or {}; wb=Workbook(); ws=wb.active; ws.title="Shape Dimensions"
    headers=["PN / MPN","Representative Ref","Manufacturer","Package / Type","Body Length (mm)","Body Width (mm)","Body Height (mm)","Overall Length (mm)","Overall Width (mm)","Paste Envelope Length (mm)","Paste Envelope Width (mm)","Paste Pad Length (mm)","Paste Pad Width (mm)","Pin / Ball Count Candidate","Pin / Ball Pitch Candidate (mm)","Pad Rows","Pad Columns","Dimension Source","Source URL","Confidence","Status","User Accepted","Remarks"]
    ws.append(headers)
    for p in unique_parts:
        r=dimension_results.get(p.mpn); s=shape_models.get(p.mpn)
        accepted=bool(getattr(s,"user_accepted",False) or getattr(r,"accepted",False))
        body_l=getattr(s,"body_length_mm",None); body_w=getattr(s,"body_width_mm",None); body_h=getattr(s,"body_height_mm",None)
        source=str(getattr(s,"source","") or "") if s else ""
        if not accepted:
            if r:
                if body_l==getattr(r,"length_mm",None): body_l=None
                if body_w==getattr(r,"width_mm",None): body_w=None
                if body_h==getattr(r,"height_mm",None): body_h=None
            elif "gerber" in source.lower():
                body_l=body_w=body_h=None
        if not s:
            body_l=getattr(r,"length_mm",None) if accepted else None
            body_w=getattr(r,"width_mm",None) if accepted else None
            body_h=getattr(r,"height_mm",None) if accepted else None
        # Paste geometry is independent measurement evidence and remains visible.
        ws.append([p.mpn,p.representative_ref,getattr(s,"manufacturer","") if s else "",getattr(s,"package_type","") if s else "",body_l,body_w,body_h,getattr(s,"overall_length_mm",None) if s else None,getattr(s,"overall_width_mm",None) if s else None,getattr(r,"paste_length_mm",None),getattr(r,"paste_width_mm",None),getattr(r,"paste_pad_length_mm",None),getattr(r,"paste_pad_width_mm",None),getattr(r,"pad_count",None),getattr(r,"pitch_mm",None),getattr(r,"pad_rows",None),getattr(r,"pad_columns",None),getattr(s,"source","") if s else getattr(r,"source",""),getattr(s,"source_url","") if s else "",getattr(s,"confidence","") if s else getattr(r,"confidence",""),getattr(s,"verification","") if s else getattr(r,"status","NOT ACCEPTED"),"YES" if accepted else "NO",getattr(s,"remarks","") if s else getattr(r,"remarks","")])
    v=wb.create_sheet("Location Verification"); v.append(["PN","Ref","CAD X (mm)","CAD Y (mm)","CAD Rotation (deg)","Gerber X (mm)","Gerber Y (mm)","Delta X (mm)","Delta Y (mm)","Status"])
    for p in unique_parts:
        c=cad_by_ref.get(p.representative_ref); r=dimension_results.get(p.mpn); gx=getattr(r,"gerber_x",None); gy=getattr(r,"gerber_y",None); cx=getattr(c,"x",None); cy=getattr(c,"y",None)
        dx=round(abs(cx-gx),4) if cx is not None and gx is not None else None; dy=round(abs(cy-gy),4) if cy is not None and gy is not None else None
        # Location verification is meaningful only for an actual Gerber match.
        # Trusted lookup/ODB dimensions without Gerber coordinates are not a
        # failed location check.
        status="PASS" if dx is not None and dy is not None and dx<=.10 and dy<=.10 else ("WARNING" if dx is not None else "NOT APPLICABLE")
        v.append([p.mpn,p.representative_ref,cx,cy,getattr(c,"rotation",None),gx,gy,dx,dy,status])
    log=wb.create_sheet("Processing Log"); log.append(["PN","Ref","Source attempted","Source selected","Accepted","Warning / Remarks","Date/time"])
    now=datetime.now().isoformat(timespec="seconds")
    for p in unique_parts:
        r=dimension_results.get(p.mpn); s=shape_models.get(p.mpn)
        selected=(getattr(s,"source","") if s else "") or getattr(r,"source","")
        accepted=bool(getattr(s,"user_accepted",False) or getattr(r,"accepted",False))
        remarks=(getattr(s,"remarks","") if s else "") or getattr(r,"remarks","")
        log.append([p.mpn,p.representative_ref,"Silkscreen; Solder Paste; MPN lookup; ODB++",selected,"YES" if accepted else "NO",remarks,now])
    for sheet in wb.worksheets:
        for cell in sheet[1]: cell.font=Font(bold=True)
        sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width=min(max(len(str(x.value or "")) for x in col)+2,42)
    wb.save(path)
