"""Step 19: acceptance/rejection gates across GUI and both export formats."""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from models import CadRecord, UniquePart
from dimensions.gerber_dimension import GerberDimensionResult
from dimensions.shape_model import build_shape_model
from export.excel_export import export_excel
from export.text_export import export_text


def _export(tmp_path, part, cad, dim, label):
    shape = build_shape_model(part, dimension=dim)
    xlsx = tmp_path / (label + ".xlsx")
    txt = tmp_path / (label + ".txt")
    export_excel(xlsx, [part], {cad.ref: cad}, {part.mpn: dim}, {part.mpn: shape})
    export_text(txt, [part], {part.mpn: dim}, {part.mpn: shape})
    ws = load_workbook(xlsx, data_only=True)["Shape Dimensions"]
    row = dict(zip([cell.value for cell in ws[1]], [cell.value for cell in ws[2]]))
    lines = txt.read_text(encoding="utf-8").splitlines()
    text_row = dict(zip(lines[0].split("\t"), lines[1].split("\t")))
    return row, text_row


def test_gui_reject_and_reaccept_are_reflected_in_both_exports(tmp_path):
    from PySide6.QtWidgets import QApplication
    from ui.main_window import MainWindow

    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    part = UniquePart("REVIEW-GATE-1", ["U1"], "U1")
    cad = CadRecord(ref="U1", x=10, y=20, rotation=0, layer="Top")
    dim = GerberDimensionResult(
        ref="U1", length_mm=4.1, width_mm=2.1,
        source="Gerber Silkscreen - Proposed", confidence="LOW",
        accepted=False, status="WAITING FOR USER ACCEPTANCE")
    window.state.unique_parts = [part]
    window.state.dimension_results = {part.mpn: dim}
    try:
        window._populate()
        window.table.selectRow(0)

        proposed_x, proposed_t = _export(tmp_path, part, cad, dim, "proposed")
        assert proposed_x["Body Length (mm)"] is None
        assert proposed_x["User Accepted"] == "NO"
        assert proposed_t["BODY_L_MM"] == ""
        assert proposed_t["USER_ACCEPTED"] == "NO"

        window._accept_dimension()
        accepted_x, accepted_t = _export(tmp_path, part, cad, dim, "accepted")
        assert accepted_x["Body Length (mm)"] == 4.1
        assert accepted_x["User Accepted"] == "YES"
        assert accepted_t["BODY_L_MM"] == "4.1"
        assert accepted_t["USER_ACCEPTED"] == "YES"

        window._reject_dimension()
        rejected_x, rejected_t = _export(tmp_path, part, cad, dim, "rejected")
        assert dim.status == "MANUAL REVIEW" and not dim.accepted
        assert rejected_x["Body Length (mm)"] is None
        assert rejected_x["User Accepted"] == "NO"
        assert rejected_t["BODY_L_MM"] == ""
        assert rejected_t["USER_ACCEPTED"] == "NO"
    finally:
        window.close()
