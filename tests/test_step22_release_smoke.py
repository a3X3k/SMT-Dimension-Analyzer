"""Step 22 release smoke: real parsers, review gate, and export round trip."""
from openpyxl import load_workbook

from parsers.cad_parser import parse_cad
from parsers.bom_parser import parse_bom, group_unique_parts
from matching.bom_cad_matcher import match_bom_to_cad
from parsers.gerber_parser import parse_gerber
from dimensions.gerber_dimension import derive_gerber_dimension
from dimensions.shape_model import build_project_shapes
from export.excel_export import export_excel
from export.text_export import export_text


def test_release_smoke_real_input_to_reviewed_outputs(tmp_path):
    cad_file = tmp_path / "placements.csv"
    cad_file.write_text(
        "Reference,X location,Y location,Angle\nU1,10,20,90\n",
        encoding="utf-8")
    bom_file = tmp_path / "bom.csv"
    bom_file.write_text("Reference,Part Number\nU1,RELEASE-1\n", encoding="utf-8")
    silk_file = tmp_path / "outline.GTO"
    silk_file.write_text(
        "%FSLAX24Y24*%\n%MOMM*%\n%ADD10C,0.100*%\nD10*\n"
        "X80000Y190000D02*X120000Y190000D01*"
        "X120000Y210000D01*X80000Y210000D01*"
        "X80000Y190000D01*\nM02*\n", encoding="utf-8")

    cad = parse_cad(cad_file)
    bom = parse_bom(bom_file)
    assert match_bom_to_cad(cad, bom)["matched"] == {"U1"}
    parts = group_unique_parts(bom)
    doc = parse_gerber(silk_file)
    dim = derive_gerber_dimension("U1", 10, 20, "Top", [doc])
    assert dim.length_mm is not None and dim.width_mm is not None
    assert not dim.accepted
    dims = {parts[0].mpn: dim}
    cad_map = {c.ref: c for c in cad}

    def snapshot(label):
        shapes = build_project_shapes(parts, cad, dims, {})
        excel = tmp_path / (label + ".xlsx")
        txt = tmp_path / (label + ".txt")
        export_excel(excel, parts, cad_map, dims, shapes)
        export_text(txt, parts, dims, shapes)
        book = load_workbook(excel, data_only=True)
        assert book.sheetnames == [
            "Shape Dimensions", "Location Verification", "Processing Log"]
        ws = book["Shape Dimensions"]
        xrow = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
        lines = txt.read_text(encoding="utf-8").splitlines()
        trow = dict(zip(lines[0].split("\t"), lines[1].split("\t")))
        return xrow, trow

    pending_x, pending_t = snapshot("pending")
    assert pending_x["Body Length (mm)"] is None
    assert pending_t["BODY_L_MM"] == ""
    assert pending_x["Gerber Length Evidence (mm)"] == dim.length_mm

    dim.accepted = True
    dim.status = "USER ACCEPTED"
    approved_x, approved_t = snapshot("approved")
    assert approved_x["Body Length (mm)"] == dim.length_mm
    assert approved_t["BODY_L_MM"] == str(dim.length_mm)
    assert approved_x["User Accepted"] == approved_t["USER_ACCEPTED"] == "YES"

    dim.accepted = False
    dim.status = "MANUAL REVIEW"
    rejected_x, rejected_t = snapshot("rejected")
    assert rejected_x["Body Length (mm)"] is None
    assert rejected_t["BODY_L_MM"] == ""
    assert rejected_x["User Accepted"] == rejected_t["USER_ACCEPTED"] == "NO"
