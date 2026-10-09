from pathlib import Path
import sqlite3
from openpyxl import Workbook, load_workbook
from parsers.bom_parser import parse_bom,group_unique_parts
from parsers.cad_parser import parse_cad
from parsers.gerber_parser import classify_gerber
from database.database import init_db
from export.excel_export import export_excel
from export.text_export import export_text

def test_bom_grouping(tmp_path):
    p=tmp_path/'bom.csv'; p.write_text('Ref,MPN\nC101,ABC123\nC102,ABC123\nR1,XYZ9\n')
    parts=group_unique_parts(parse_bom(p)); assert len(parts)==2; assert parts[0].representative_ref=='C101'; assert parts[0].refs==['C101','C102']

def test_cad_parse(tmp_path):
    p=tmp_path/'cad.csv'; p.write_text('Reference,X,Y,Rotation,Layer\nC101,1.2,3.4,90,Top\n')
    r=parse_cad(p)[0]; assert r.ref=='C101' and r.x==1.2 and r.y==3.4 and r.rotation==90 and r.layer=='Top'

def test_layer_classification():
    assert classify_gerber('a.GTO')=='Top Silkscreen'
    assert classify_gerber('a.GTP')=='Top Paste'

def test_db_init(tmp_path):
    p=init_db(tmp_path/'x.db')
    with sqlite3.connect(p) as con:
        names={r[0] for r in con.execute("select name from sqlite_master where type='table'")}
    assert {'components','projects','component_results'}<=names

def test_export_framework(tmp_path):
    bom=tmp_path/'b.csv'; bom.write_text('Ref,MPN\nC1,P1\n')
    parts=group_unique_parts(parse_bom(bom))
    x=tmp_path/'Shape_Dimensions.xlsx'; t=tmp_path/'Shape_Dimensions.txt'
    export_excel(x,parts); export_text(t,parts)
    wb=load_workbook(x); assert wb.sheetnames==['Shape Dimensions','Location Verification','Processing Log']
    text=t.read_text()
    rows=text.splitlines()
    assert rows[0].split(chr(9))[:2]==['MPN','REF']
    assert rows[1].split(chr(9))[:2]==['P1','C1']


def test_cad_excel_finds_header_after_report_metadata(tmp_path):
    from openpyxl import Workbook
    p=tmp_path/'placement.xlsx'; wb=Workbook(); ws=wb.active
    ws.append(['PCB Placement Report']); ws.append(['Generated','today']); ws.append([])
    ws.append(['Component Reference','Centre X','Centre Y','Orientation','PCB Side'])
    ws.append(['U1','12.5 mm','7.25 mm','90 deg','Top']); wb.save(p)
    r=parse_cad(p)[0]
    assert r.ref=='U1' and r.x==12.5 and r.y==7.25 and r.rotation==90 and r.layer=='Top'


def test_cad_mils_conversion_and_optional_part_number(tmp_path):
    from parsers.cad_parser import parse_cad
    path=tmp_path/"placement.txt"
    path.write_text("refdes\tsymbol_x\tsymbol_y\trotation\tmirror\nR1\t1000\t-500\t90\tTop\n")
    records=parse_cad(path,units="mils")
    assert len(records)==1
    assert records[0].ref=="R1"
    assert abs(records[0].x-25.4)<1e-9
    assert abs(records[0].y+12.7)<1e-9
    assert records[0].rotation==90
    assert records[0].mpn==""
    assert records[0].layer=="Top"


def test_cad_excel_header_aliases_and_mirror_boolean(tmp_path):
    from parsers.cad_parser import parse_cad
    import pandas as pd
    path=tmp_path/"placement.xlsx"
    pd.DataFrame([{"REFDES":"U1","SYM_X":12.5,"SYM_Y":4.2,"SYM_ROTATE":270,"SYM_MIRROR":"NO"}]).to_excel(path,index=False)
    records=parse_cad(path,units="mm")
    assert len(records)==1
    assert records[0].x==12.5 and records[0].y==4.2
    assert records[0].rotation==270
    assert records[0].layer==""
