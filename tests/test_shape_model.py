from models import UniquePart, CadRecord
from dimensions.gerber_dimension import GerberDimensionResult
from dimensions.shape_model import build_shape_model, recognize_package_family
from lookup.mpn_lookup import MpnData

def test_package_family_recognition():
    assert recognize_package_family("VQFN-32")[0] == "QFN/DFN"
    assert recognize_package_family("BGA")[0] == "BGA"
    assert recognize_package_family("")[0] == "UNKNOWN"

def test_gerber_paste_never_becomes_physical_pin_data():
    part=UniquePart("ABC",["U1"],"U1")
    dim=GerberDimensionResult(ref="U1",length_mm=4.0,width_mm=4.0,pad_count=32,pitch_mm=0.5,pad_rows=4,pad_columns=8)
    lookup=MpnData(query="ABC",package_type="VQFN-32",status="NOT FOUND")
    s=build_shape_model(part,dimension=dim,lookup=lookup)
    assert s.pin_count is None and s.pin_pitch_mm is None
    assert s.bga_rows is None and s.bga_columns is None and s.ball_pitch_mm is None

def test_exact_lookup_physical_data_has_priority():
    part=UniquePart("ABC",["U1"],"U1")
    dim=GerberDimensionResult(ref="U1",length_mm=5.0,width_mm=5.0,pad_count=40,pitch_mm=0.4)
    lookup=MpnData(query="ABC",matched_mpn="ABC",package_type="QFN",body_length_mm=4.0,body_width_mm=4.0,pin_count=32,pin_pitch_mm=0.5,source="Manufacturer",confidence="HIGH",status="EXACT MPN MATCH")
    s=build_shape_model(part,dimension=dim,lookup=lookup)
    assert (s.body_length_mm,s.body_width_mm)==(4.0,4.0)
    assert (s.pin_count,s.pin_pitch_mm)==(32,0.5)

def test_exact_lookup_without_dimensions_does_not_claim_dimension_provenance():
    part=UniquePart("ABC",["U1"],"U1")
    lookup=MpnData(query="ABC",matched_mpn="ABC",manufacturer="Acme",package_type="QFN",source="DigiKey",source_url="https://example.invalid/product",confidence="HIGH",status="EXACT MPN MATCH")
    s=build_shape_model(part,lookup=lookup)
    assert s.manufacturer=="Acme" and s.package_type=="QFN"
    assert s.body_length_mm is None and s.body_width_mm is None
    assert s.source=="" and s.source_url=="" and s.confidence=="NONE"

def test_exact_lookup_metadata_does_not_mask_gerber_dimension_source():
    part=UniquePart("ABC",["U1"],"U1")
    lookup=MpnData(query="ABC",matched_mpn="ABC",package_type="QFN",source="DigiKey",confidence="HIGH",status="EXACT MPN MATCH")
    dim=GerberDimensionResult(ref="U1",length_mm=5.0,width_mm=4.0,source="Gerber Silkscreen - Proposed",confidence="MEDIUM")
    s=build_shape_model(part,dimension=dim,lookup=lookup)
    assert (s.body_length_mm,s.body_width_mm)==(5.0,4.0)
    assert s.source=="Gerber Silkscreen - Proposed" and s.confidence=="MEDIUM"

def test_txt_export_hides_unaccepted_gerber_body_but_keeps_candidate(tmp_path):
    from export.text_export import export_text
    part=UniquePart("ABC",["U1"],"U1")
    dim=GerberDimensionResult(ref="U1",length_mm=5.0,width_mm=4.0,source="Gerber Silkscreen - Proposed",confidence="MEDIUM",accepted=False)
    s=build_shape_model(part,dimension=dim)
    out=tmp_path/"out.txt"
    export_text(out,[part],{"ABC":dim},{"ABC":s})
    headers,values=out.read_text().splitlines()
    row=dict(zip(headers.split("\t"),values.split("\t")))
    assert row["BODY_L_MM"]=="" and row["BODY_W_MM"]==""
    assert row["GERBER_L_MM"]=="5.0" and row["GERBER_W_MM"]=="4.0"
    assert row["USER_ACCEPTED"]=="NO"

def test_txt_export_promotes_accepted_gerber_body(tmp_path):
    from export.text_export import export_text
    part=UniquePart("ABC",["U1"],"U1")
    dim=GerberDimensionResult(ref="U1",length_mm=5.0,width_mm=4.0,source="Gerber Silkscreen - Proposed",confidence="MEDIUM",accepted=True)
    s=build_shape_model(part,dimension=dim)
    out=tmp_path/"out.txt"
    export_text(out,[part],{"ABC":dim},{"ABC":s})
    headers,values=out.read_text().splitlines()
    row=dict(zip(headers.split("\t"),values.split("\t")))
    assert row["BODY_L_MM"]=="5.0" and row["BODY_W_MM"]=="4.0"
    assert row["USER_ACCEPTED"]=="YES"


def test_review_ui_labels_distinguish_paste_candidates_and_provenance():
    from pathlib import Path
    source=Path("ui/main_window.py").read_text(encoding="utf-8")
    assert "Paste Envelope L×W (mm)" in source
    assert "Paste Pad L×W / Count / Pitch (mm)" in source
    assert "Paste geometry only — not body size" in source
    assert "Dimension Source / Confidence" in source
    assert "Review Status" in source
    assert "Confidence {getattr(x,'confidence','') or 'NONE'}" in source
    assert "Accepted {'YES' if getattr(x,'accepted',False) else 'NO'}" in source


def test_exports_keep_unaccepted_paste_but_gate_body_dimensions(tmp_path):
    from models import UniquePart
    from dimensions.gerber_dimension import GerberDimensionResult
    from dimensions.shape_model import build_shape_model
    from export.excel_export import export_excel
    from export.text_export import export_text
    from openpyxl import load_workbook

    part=UniquePart('MPN1',['U1'],'U1')
    result=GerberDimensionResult(
        'U1',length_mm=4.0,width_mm=2.0,
        source='Gerber Silkscreen - Proposed',confidence='MEDIUM',
        paste_length_mm=5.2,paste_width_mm=2.6,
        paste_pad_length_mm=1.0,paste_pad_width_mm=.5,
        accepted=False)
    shape=build_shape_model(part,dimension=result)

    xlsx=tmp_path/'out.xlsx'
    txt=tmp_path/'out.txt'
    export_excel(xlsx,[part],{}, {'MPN1':result})
    export_text(txt,[part],{'MPN1':result},{'MPN1':shape})

    ws=load_workbook(xlsx,data_only=True)['Shape Dimensions']
    headers=[c.value for c in ws[1]]
    row=[c.value for c in ws[2]]
    data=dict(zip(headers,row))
    assert data['Body Length (mm)'] is None
    assert data['Body Width (mm)'] is None
    assert data['Paste Envelope Length (mm)']==5.2
    assert data['Paste Pad Length (mm)']==1.0

    lines=txt.read_text(encoding='utf-8').splitlines()
    data=dict(zip(lines[0].split('\t'),lines[1].split('\t')))
    assert data['BODY_L_MM']==''
    assert data['BODY_W_MM']==''
    assert data['PASTE_ENVELOPE_L_MM']=='5.2'
    assert data['PASTE_PAD_L_MM']=='1.0'
    assert data['USER_ACCEPTED']=='NO'


def test_exact_mpn_structured_metric_dimension_extraction_is_conservative():
    from lookup.mpn_lookup import _exact_metric_dimensions
    product={'Parameters':[
        {'ParameterText':'Length - Overall','ValueText':'4.90 mm'},
        {'ParameterText':'Width - Overall','ValueText':'3.90mm'},
        {'ParameterText':'Height - Seated (Max)','ValueText':'1.75 mm'},
        {'ParameterText':'Body Length','ValueText':'0.25 in'},
    ]}
    d=_exact_metric_dimensions(product)
    assert d=={'body_length_mm':4.9,'body_width_mm':3.9,'body_height_mm':1.75}

def test_exact_mpn_dimension_extraction_does_not_guess_units():
    from lookup.mpn_lookup import _exact_metric_dimensions
    product={'Parameters':[
        {'ParameterText':'Body Length','ValueText':'0.25 in'},
        {'ParameterText':'Body Width','ValueText':'250 mil'},
        {'ParameterText':'Body Height','ValueText':'unknown'},
        {'ParameterText':'Length - Overall','ValueText':'4.9'},
    ]}
    assert _exact_metric_dimensions(product)=={}


def test_review_lookup_uses_expanded_source_and_status_columns():
    from pathlib import Path
    source=Path('ui/main_window.py').read_text(encoding='utf-8')
    assert 'self.table.setItem(r,9,QTableWidgetItem(data.source or "MPN lookup"))' in source
    assert 'self.table.setItem(r,10,QTableWidgetItem(data.status))' in source
    assert 'self.table.setItem(r,8,QTableWidgetItem(data.source or "MPN lookup"))' not in source


def test_lookup_dimensions_are_not_marked_user_accepted_by_unrelated_gerber():
    from types import SimpleNamespace
    from models import UniquePart
    from dimensions.shape_model import build_shape_model
    from dimensions.gerber_dimension import GerberDimensionResult
    part=UniquePart('ABC',['U1'],'U1')
    lookup=SimpleNamespace(status='EXACT MPN MATCH',body_length_mm=5.0,body_width_mm=4.0,
        body_height_mm=1.0,pin_count=None,pin_pitch_mm=None,lead_width_mm=None,lead_length_mm=None,
        bga_rows=None,bga_columns=None,ball_pitch_mm=None,source='Mouser',datasheet_url='',source_url='',confidence='HIGH',
        package_type='',manufacturer='')
    gerber=GerberDimensionResult('U1',length_mm=6.0,width_mm=6.0,accepted=True,status='USER ACCEPTED')
    shape=build_shape_model(part,dimension=gerber,lookup=lookup)
    assert shape.body_length_mm==5.0
    assert shape.body_width_mm==4.0
    assert shape.user_accepted is False
    assert shape.verification=='NOT REQUIRED'


def test_text_export_gates_only_unaccepted_gerber_fields_in_mixed_source(tmp_path):
    from models import UniquePart
    from dimensions.shape_model import ShapeModel
    from dimensions.gerber_dimension import GerberDimensionResult
    from export.text_export import export_text
    part=UniquePart('ABC',['U1'],'U1')
    shape=ShapeModel(mpn='ABC',ref='U1',body_length_mm=5.0,body_width_mm=2.0,
        body_height_mm=1.0,source='Mouser + Gerber Silkscreen - Proposed',confidence='HIGH')
    gerber=GerberDimensionResult('U1',length_mm=5.0,width_mm=2.0,height_mm=None,accepted=False)
    out=tmp_path/'mixed.txt'
    export_text(out,[part],{'ABC':gerber},{'ABC':shape})
    lines=out.read_text(encoding='utf-8').splitlines()
    data=dict(zip(lines[0].split('\t'),lines[1].split('\t')))
    assert data['BODY_L_MM']==''
    assert data['BODY_W_MM']==''
    assert data['BODY_H_MM']=='1.0'


def test_text_export_gates_unaccepted_gerber_shape_without_raw_result(tmp_path):
    from models import UniquePart
    from dimensions.shape_model import ShapeModel
    from export.text_export import export_text
    part=UniquePart('ABC',['U1'],'U1')
    shape=ShapeModel(mpn='ABC',ref='U1',body_length_mm=4.0,body_width_mm=2.0,
        source='Gerber Silkscreen - Proposed',confidence='MEDIUM',user_accepted=False)
    out=tmp_path/'shape_only.txt'
    export_text(out,[part],{}, {'ABC':shape})
    lines=out.read_text(encoding='utf-8').splitlines()
    data=dict(zip(lines[0].split('\t'),lines[1].split('\t')))
    assert data['BODY_L_MM']==''
    assert data['BODY_W_MM']==''


def test_excel_export_preserves_trusted_shape_dimensions_and_gates_gerber(tmp_path):
    from openpyxl import load_workbook
    from models import UniquePart
    from dimensions.shape_model import ShapeModel
    from dimensions.gerber_dimension import GerberDimensionResult
    from export.excel_export import export_excel
    part=UniquePart('ABC',['U1'],'U1')
    shape=ShapeModel(mpn='ABC',ref='U1',manufacturer='Acme',body_length_mm=5.0,
        body_width_mm=2.0,body_height_mm=1.0,source='Mouser + Gerber Silkscreen - Proposed',confidence='HIGH')
    gerber=GerberDimensionResult('U1',length_mm=5.0,width_mm=2.0,accepted=False,
        paste_length_mm=6.0,paste_width_mm=3.0)
    out=tmp_path/'mixed.xlsx'
    export_excel(out,[part],{}, {'ABC':gerber},{'ABC':shape})
    ws=load_workbook(out,data_only=True)['Shape Dimensions']
    data=dict(zip([x.value for x in ws[1]],[x.value for x in ws[2]]))
    assert data['Body Length (mm)'] is None
    assert data['Body Width (mm)'] is None
    assert data['Body Height (mm)']==1.0
    assert data['Paste Envelope Length (mm)']==6.0
    assert data['Manufacturer']=='Acme'


def test_excel_processing_log_uses_selected_shape_provenance(tmp_path):
    from openpyxl import load_workbook
    from models import UniquePart
    from dimensions.shape_model import ShapeModel
    from export.excel_export import export_excel
    part=UniquePart('ABC',['U1'],'U1')
    shape=ShapeModel(mpn='ABC',ref='U1',body_length_mm=5.0,source='Mouser',confidence='HIGH')
    out=tmp_path/'log.xlsx'
    export_excel(out,[part],{}, {}, {'ABC':shape})
    ws=load_workbook(out,data_only=True)['Processing Log']
    data=dict(zip([x.value for x in ws[1]],[x.value for x in ws[2]]))
    assert data['Source selected']=='Mouser'
    assert data['Review Required']=='NO'
    assert data['Review Accepted']=='NOT REQUIRED'
    assert 'MPN lookup' in data['Source attempted']


def test_excel_location_verification_is_not_applicable_without_gerber_match(tmp_path):
    from openpyxl import load_workbook
    from models import UniquePart
    from export.excel_export import export_excel
    part=UniquePart('ABC',['U1'],'U1')
    out=tmp_path/'location.xlsx'
    export_excel(out,[part],{}, {}, {})
    ws=load_workbook(out,data_only=True)['Location Verification']
    data=dict(zip([x.value for x in ws[1]],[x.value for x in ws[2]]))
    assert data['Status']=='NOT APPLICABLE'


def test_excel_location_pass_does_not_claim_rotation_verified(tmp_path):
    from openpyxl import load_workbook
    from types import SimpleNamespace
    from models import UniquePart
    from dimensions.gerber_dimension import GerberDimensionResult
    from export.excel_export import export_excel
    part=UniquePart('ABC',['U1'],'U1')
    cad=SimpleNamespace(x=10.0,y=20.0,rotation=90.0)
    gerber=GerberDimensionResult('U1',gerber_x=10.05,gerber_y=19.95)
    out=tmp_path/'rotation.xlsx'
    export_excel(out,[part],{'U1':cad},{'ABC':gerber},{})
    ws=load_workbook(out,data_only=True)['Location Verification']
    data=dict(zip([x.value for x in ws[1]],[x.value for x in ws[2]]))
    assert data['Rotation Check']=='NOT AVAILABLE'
    assert data['Status']=='POSITION PASS / ROTATION NOT VERIFIED'


def test_per_field_provenance_preserves_trusted_value_equal_to_gerber(tmp_path):
    from openpyxl import load_workbook
    from models import UniquePart
    from dimensions.shape_model import ShapeModel
    from dimensions.gerber_dimension import GerberDimensionResult
    from export.excel_export import export_excel
    from export.text_export import export_text
    part=UniquePart('ABC',['U1'],'U1')
    shape=ShapeModel(mpn='ABC',ref='U1',body_length_mm=5.0,body_width_mm=2.0,body_height_mm=1.0,
        body_length_source='Mouser',body_width_source='Gerber Silkscreen - Proposed',body_height_source='Mouser',
        source='Mouser + Mouser + Gerber Silkscreen - Proposed',confidence='HIGH')
    gerber=GerberDimensionResult('U1',length_mm=5.0,width_mm=2.0,accepted=False)
    txt=tmp_path/'provenance.txt'
    export_text(txt,[part],{'ABC':gerber},{'ABC':shape})
    lines=txt.read_text(encoding='utf-8').splitlines()
    data=dict(zip(lines[0].split('\t'),lines[1].split('\t')))
    assert data['BODY_L_MM']=='5.0'
    assert data['BODY_W_MM']==''
    assert data['BODY_H_MM']=='1.0'
    xlsx=tmp_path/'provenance.xlsx'
    export_excel(xlsx,[part],{}, {'ABC':gerber},{'ABC':shape})
    ws=load_workbook(xlsx,data_only=True)['Shape Dimensions']
    row=dict(zip([x.value for x in ws[1]],[x.value for x in ws[2]]))
    assert row['Body Length (mm)']==5.0
    assert row['Body Width (mm)'] is None
    assert row['Body Height (mm)']==1.0


def test_text_export_keeps_unaccepted_gerber_height_as_evidence(tmp_path):
    from models import UniquePart
    from dimensions.gerber_dimension import GerberDimensionResult
    from export.text_export import export_text
    part=UniquePart('ABC',['U1'],'U1')
    gerber=GerberDimensionResult('U1',length_mm=5.0,width_mm=2.0,height_mm=1.2,accepted=False)
    out=tmp_path/'height.txt'
    export_text(out,[part],{'ABC':gerber},{})
    lines=out.read_text(encoding='utf-8').splitlines()
    data=dict(zip(lines[0].split('\t'),lines[1].split('\t')))
    assert data['BODY_H_MM']==''
    assert data['GERBER_H_MM']=='1.2'


def test_excel_separates_unaccepted_gerber_body_evidence(tmp_path):
    from openpyxl import load_workbook
    from models import UniquePart
    from dimensions.gerber_dimension import GerberDimensionResult
    from export.excel_export import export_excel
    part=UniquePart('ABC',['U1'],'U1')
    gerber=GerberDimensionResult('U1',length_mm=5.0,width_mm=2.0,height_mm=1.2,accepted=False)
    out=tmp_path/'evidence.xlsx'
    export_excel(out,[part],{}, {'ABC':gerber},{})
    ws=load_workbook(out,data_only=True)['Shape Dimensions']
    data=dict(zip([x.value for x in ws[1]],[x.value for x in ws[2]]))
    assert data['Body Length (mm)'] is None
    assert data['Body Width (mm)'] is None
    assert data['Body Height (mm)'] is None
    assert data['Gerber Length Evidence (mm)']==5.0
    assert data['Gerber Width Evidence (mm)']==2.0
    assert data['Gerber Height Evidence (mm)']==1.2


def test_exports_surface_per_field_body_provenance(tmp_path):
    from openpyxl import load_workbook
    from models import UniquePart
    from dimensions.shape_model import ShapeModel
    from export.excel_export import export_excel
    from export.text_export import export_text
    part=UniquePart('ABC',['U1'],'U1')
    shape=ShapeModel(mpn='ABC',ref='U1',body_length_mm=5.0,body_width_mm=2.0,body_height_mm=1.0,
        body_length_source='Mouser',body_width_source='Mouser',body_height_source='Gerber Silkscreen - Proposed',
        source='Mouser + Mouser + Gerber Silkscreen - Proposed')
    txt=tmp_path/'sources.txt'
    export_text(txt,[part],{}, {'ABC':shape})
    lines=txt.read_text(encoding='utf-8').splitlines()
    data=dict(zip(lines[0].split('\t'),lines[1].split('\t')))
    assert data['BODY_L_SOURCE']=='Mouser'
    assert data['BODY_W_SOURCE']=='Mouser'
    assert data['BODY_H_SOURCE']=='Gerber Silkscreen - Proposed'
    xlsx=tmp_path/'sources.xlsx'
    export_excel(xlsx,[part],{}, {}, {'ABC':shape})
    ws=load_workbook(xlsx,data_only=True)['Shape Dimensions']
    row=dict(zip([x.value for x in ws[1]],[x.value for x in ws[2]]))
    assert row['Body Length Source']=='Mouser'
    assert row['Body Width Source']=='Mouser'
    assert row['Body Height Source']=='Gerber Silkscreen - Proposed'


def test_processing_log_does_not_call_trusted_dimensions_rejected(tmp_path):
    from openpyxl import load_workbook
    from models import UniquePart
    from dimensions.shape_model import ShapeModel
    from export.excel_export import export_excel
    part=UniquePart('ABC',['U1'],'U1')
    shape=ShapeModel(mpn='ABC',ref='U1',body_length_mm=5.0,body_width_mm=2.0,
        body_length_source='Mouser',body_width_source='Mouser',source='Mouser',
        verification='NOT REQUIRED',user_accepted=False)
    out=tmp_path/'trusted.xlsx'
    export_excel(out,[part],{}, {}, {'ABC':shape})
    ws=load_workbook(out,data_only=True)['Processing Log']
    data=dict(zip([x.value for x in ws[1]],[x.value for x in ws[2]]))
    assert data['Review Required']=='NO'
    assert data['Review Accepted']=='NOT REQUIRED'


def test_processing_log_marks_unaccepted_gerber_as_review_required(tmp_path):
    from openpyxl import load_workbook
    from models import UniquePart
    from dimensions.shape_model import ShapeModel
    from dimensions.gerber_dimension import GerberDimensionResult
    from export.excel_export import export_excel
    part=UniquePart('ABC',['U1'],'U1')
    shape=ShapeModel(mpn='ABC',ref='U1',body_length_mm=5.0,
        body_length_source='Gerber Silkscreen - Proposed',
        source='Gerber Silkscreen - Proposed',user_accepted=False)
    gerber=GerberDimensionResult('U1',length_mm=5.0,accepted=False)
    out=tmp_path/'review.xlsx'
    export_excel(out,[part],{}, {'ABC':gerber},{'ABC':shape})
    ws=load_workbook(out,data_only=True)['Processing Log']
    data=dict(zip([x.value for x in ws[1]],[x.value for x in ws[2]]))
    assert data['Review Required']=='YES'
    assert data['Review Accepted']=='NO'


def test_manual_body_override_survives_shape_rebuild_and_export(tmp_path):
    from models import UniquePart
    from dimensions.gerber_dimension import GerberDimensionResult
    from dimensions.shape_model import build_shape_model
    from export.text_export import export_text
    part=UniquePart('ABC',['U1'],'U1')
    edited=GerberDimensionResult('U1',length_mm=5.5,width_mm=2.2,source='USER - Manual Body Adjustment',
        confidence='USER CONFIRMED',status='WAITING FOR USER ACCEPTANCE',accepted=False)
    edited.manual_body_override=True
    shape=build_shape_model(part,None,edited,None)
    assert shape.body_length_source=='USER'
    assert shape.body_width_source=='USER'
    assert shape.verification=='USER REVIEW'
    out=tmp_path/'manual.txt'
    export_text(out,[part],{'ABC':edited},{'ABC':shape})
    lines=out.read_text(encoding='utf-8').splitlines()
    data=dict(zip(lines[0].split('\t'),lines[1].split('\t')))
    assert data['BODY_L_MM']=='5.5'
    assert data['BODY_W_MM']=='2.2'
    assert data['BODY_L_SOURCE']=='USER'
    assert data['BODY_W_SOURCE']=='USER'
