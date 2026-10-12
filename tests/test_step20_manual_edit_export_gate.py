"""Step 20: a manual edit invalidates prior acceptance in both exports."""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from models import CadRecord, UniquePart
from dimensions.gerber_dimension import GerberDimensionResult
from dimensions.shape_model import build_shape_model
from export.excel_export import export_excel
from export.text_export import export_text


def _snapshot(tmp_path, part, cad, dim, label):
    shape = build_shape_model(part, dimension=dim)
    excel = tmp_path / (label + ".xlsx")
    text = tmp_path / (label + ".txt")
    export_excel(excel, [part], {cad.ref: cad}, {part.mpn: dim}, {part.mpn: shape})
    export_text(text, [part], {part.mpn: dim}, {part.mpn: shape})
    ws = load_workbook(excel, data_only=True)["Shape Dimensions"]
    x = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
    lines = text.read_text(encoding="utf-8").splitlines()
    t = dict(zip(lines[0].split("\t"), lines[1].split("\t")))
    return x, t


def test_manual_edit_revokes_acceptance_until_explicit_reaccept(tmp_path):
    from PySide6.QtWidgets import QApplication
    from ui.main_window import MainWindow

    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    part = UniquePart("EDIT-GATE-1", ["U1"], "U1")
    cad = CadRecord(ref="U1", x=10, y=20, rotation=0, layer="Top")
    dim = GerberDimensionResult(
        ref="U1", length_mm=4.0, width_mm=2.0,
        source="Gerber Silkscreen - Proposed", accepted=False)
    window.state.unique_parts = [part]
    window.state.dimension_results = {part.mpn: dim}
    try:
        window._populate()
        window.table.selectRow(0)
        window._accept_dimension()
        assert dim.accepted
        before_x, before_t = _snapshot(tmp_path, part, cad, dim, "before")
        assert before_x["Body Length (mm)"] == 4.0
        assert before_t["BODY_L_MM"] == "4.0"

        window.manual_l.setValue(5.5)
        window.manual_w.setValue(2.5)
        window._apply_adjust()
        assert dim.manual_body_override
        assert dim.status == "WAITING FOR USER ACCEPTANCE"
        assert dim.accepted is False
        pending_x, pending_t = _snapshot(tmp_path, part, cad, dim, "pending")
        assert pending_x["Body Length (mm)"] is None
        assert pending_x["User Accepted"] == "NO"
        assert pending_t["BODY_L_MM"] == ""
        assert pending_t["USER_ACCEPTED"] == "NO"

        window._accept_dimension()
        accepted_x, accepted_t = _snapshot(tmp_path, part, cad, dim, "reaccepted")
        assert accepted_x["Body Length (mm)"] == 5.5
        assert accepted_x["Body Width (mm)"] == 2.5
        assert accepted_x["User Accepted"] == "YES"
        assert accepted_t["BODY_L_MM"] == "5.5"
        assert accepted_t["USER_ACCEPTED"] == "YES"
    finally:
        window.close()
