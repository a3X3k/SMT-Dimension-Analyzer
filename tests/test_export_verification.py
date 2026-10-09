from openpyxl import load_workbook
from models import UniquePart, CadRecord
from dimensions.gerber_dimension import GerberDimensionResult
from dimensions.shape_model import build_shape_model
from export.excel_export import export_excel
from export.text_export import export_text


def _rows(path):
    lines=path.read_text(encoding="utf-8").splitlines()
    return dict(zip(lines[0].split("\t"),lines[1].split("\t")))


def _excel_row(ws):
    return dict(zip([c.value for c in ws[1]],[c.value for c in ws[2]]))


def test_export_review_provenance_and_location_tolerance(tmp_path):
    part=UniquePart("MPN-1",["U1"],"U1")
    cad={"U1":CadRecord(ref="U1",x=25.4,y=12.7,rotation=90)}
    dim=GerberDimensionResult(ref="U1",length_mm=4.0,width_mm=2.0,
        gerber_x=25.46,gerber_y=12.76,
        source="Gerber Silkscreen - Proposed",confidence="MEDIUM",
        paste_length_mm=4.8,paste_width_mm=2.5,
        remarks="Needs body confirmation",accepted=False)
    shape=build_shape_model(part,dimension=dim)
    x=tmp_path/"result.xlsx";t=tmp_path/"result.txt"
    export_excel(x,[part],cad,{"MPN-1":dim},{"MPN-1":shape})
    export_text(t,[part],{"MPN-1":dim},{"MPN-1":shape})
    wb=load_workbook(x,data_only=True)
    assert wb.sheetnames==["Shape Dimensions","Location Verification","Processing Log"]
    main=_excel_row(wb["Shape Dimensions"])
    assert main["Body Length (mm)"] is None
    assert main["Gerber Length Evidence (mm)"]==4.0
    assert main["Paste Envelope Length (mm)"]==4.8
    assert main["User Accepted"]=="NO"
    loc=_excel_row(wb["Location Verification"])
    assert loc["CAD X (mm)"]==25.4
    assert loc["Delta X (mm)"]==0.06
    assert loc["Delta Y (mm)"]==0.06
    assert loc["Rotation Check"]=="NOT AVAILABLE"
    assert loc["Status"]=="POSITION PASS / ROTATION NOT VERIFIED"
    log=_excel_row(wb["Processing Log"])
    assert log["Review Required"]=="YES"
    assert log["Review Accepted"]=="NO"
    assert "Needs body confirmation" in (log["Warning / Remarks"] or "")
    assert log["Date/time"]
    txt=_rows(t)
    assert txt["BODY_L_MM"]==""
    assert txt["GERBER_L_MM"]=="4.0"
    assert txt["PASTE_ENVELOPE_L_MM"]=="4.8"
    assert txt["USER_ACCEPTED"]=="NO"
    assert "Needs body confirmation" in txt["REMARKS"]


def test_export_accepted_body_and_location_warning(tmp_path):
    part=UniquePart("MPN-2",["R1"],"R1")
    cad={"R1":CadRecord(ref="R1",x=10,y=20,rotation=0)}
    dim=GerberDimensionResult(ref="R1",length_mm=3.2,width_mm=1.6,
        gerber_x=10.15,gerber_y=20,
        source="Gerber Silkscreen - Proposed",accepted=True)
    shape=build_shape_model(part,dimension=dim)
    x=tmp_path/"accepted.xlsx";t=tmp_path/"accepted.txt"
    export_excel(x,[part],cad,{"MPN-2":dim},{"MPN-2":shape})
    export_text(t,[part],{"MPN-2":dim},{"MPN-2":shape})
    wb=load_workbook(x,data_only=True)
    main=_excel_row(wb["Shape Dimensions"])
    assert main["Body Length (mm)"]==3.2
    assert main["Body Width (mm)"]==1.6
    assert main["User Accepted"]=="YES"
    assert _excel_row(wb["Location Verification"])["Status"]=="WARNING"
    assert _excel_row(wb["Processing Log"])["Review Accepted"]=="YES"
    assert _rows(t)["BODY_L_MM"]=="3.2"
    assert _rows(t)["USER_ACCEPTED"]=="YES"
