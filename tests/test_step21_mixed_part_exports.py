"""Step 21: mixed-source multi-part export isolation and review gates."""
from openpyxl import load_workbook

from models import CadRecord, UniquePart
from dimensions.gerber_dimension import GerberDimensionResult
from dimensions.shape_model import build_shape_model
from export.excel_export import export_excel
from export.text_export import export_text
from lookup.mpn_lookup import MpnData


def test_multi_part_exports_isolate_acceptance_and_preserve_lookup(tmp_path):
    parts = [
        UniquePart("ACCEPTED", ["U1"], "U1"),
        UniquePart("REJECTED", ["U2"], "U2"),
        UniquePart("LOOKUP", ["U3"], "U3"),
        UniquePart("MANUAL-PENDING", ["U4"], "U4"),
    ]
    cad = {p.representative_ref: CadRecord(ref=p.representative_ref, x=i * 10, y=20)
           for i, p in enumerate(parts, start=1)}
    dims = {
        "ACCEPTED": GerberDimensionResult(
            "U1", length_mm=4.0, width_mm=2.0,
            source="Gerber Silkscreen - Proposed", accepted=True),
        "REJECTED": GerberDimensionResult(
            "U2", length_mm=6.0, width_mm=3.0,
            source="Gerber Silkscreen - Proposed", status="MANUAL REVIEW", accepted=False),
        "MANUAL-PENDING": GerberDimensionResult(
            "U4", length_mm=7.0, width_mm=3.5,
            source="USER - Manual Body Adjustment",
            status="WAITING FOR USER ACCEPTANCE", accepted=False),
    }
    dims["MANUAL-PENDING"].manual_body_override = True
    lookup = MpnData(
        query="LOOKUP", matched_mpn="LOOKUP", status="EXACT MPN MATCH",
        body_length_mm=8.0, body_width_mm=4.0,
        source="Manufacturer", confidence="HIGH")
    shapes = {
        p.mpn: build_shape_model(
            p, dimension=dims.get(p.mpn),
            lookup=lookup if p.mpn == "LOOKUP" else None)
        for p in parts
    }

    xlsx = tmp_path / "mixed.xlsx"
    txt = tmp_path / "mixed.txt"
    export_excel(xlsx, parts, cad, dims, shapes)
    export_text(txt, parts, dims, shapes)

    ws = load_workbook(xlsx, data_only=True)["Shape Dimensions"]
    headers = [cell.value for cell in ws[1]]
    excel = {row[0]: dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)}
    lines = txt.read_text(encoding="utf-8").splitlines()
    text_headers = lines[0].split("\t")
    text = {row[0]: dict(zip(text_headers, row)) for row in
            (line.split("\t") for line in lines[1:])}

    assert len(excel) == len(text) == 4
    for mpn, expected_length, accepted in (
        ("ACCEPTED", 4.0, "YES"),
        ("REJECTED", None, "NO"),
        ("LOOKUP", 8.0, "NO"),
        ("MANUAL-PENDING", None, "NO"),
    ):
        assert excel[mpn]["Body Length (mm)"] == expected_length
        assert excel[mpn]["User Accepted"] == accepted
        assert text[mpn]["BODY_L_MM"] == (str(expected_length) if expected_length is not None else "")
        assert text[mpn]["USER_ACCEPTED"] == accepted

    assert excel["REJECTED"]["Gerber Length Evidence (mm)"] == 6.0
    assert excel["MANUAL-PENDING"]["Body Length Source"] == "USER"
    assert excel["LOOKUP"]["Body Length Source"] == "Manufacturer"
    assert text["LOOKUP"]["BODY_L_SOURCE"] == "Manufacturer"
