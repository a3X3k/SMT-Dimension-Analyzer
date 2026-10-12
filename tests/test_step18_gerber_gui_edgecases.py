"""Step 18: synthetic Gerber edge cases and review-to-export safety."""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from models import CadRecord, UniquePart
from parsers.gerber_parser import parse_gerber
from dimensions.gerber_dimension import derive_gerber_dimension
from dimensions.shape_model import build_shape_model
from export.excel_export import export_excel


def _gerber(tmp_path, filename, commands):
    path = tmp_path / filename
    path.write_text("%FSLAX24Y24*%\n%MOMM*%\n%ADD10C,0.100*%\nD10*\n" + commands + "\nM02*\n", encoding="utf-8")
    return parse_gerber(path)


def _rectangle():
    return ("X80000Y190000D02*X120000Y190000D01*"
            "X120000Y210000D01*X80000Y210000D01*"
            "X80000Y190000D01*")


def test_incomplete_single_edge_cannot_propose_body(tmp_path):
    doc = _gerber(tmp_path, "incomplete.GTO", "X80000Y190000D02*X120000Y190000D01*")
    result = derive_gerber_dimension("U1", 10, 20, "Top", [doc])
    assert result.length_mm is None
    assert result.accepted is False
    assert result.status == "NOT AVAILABLE / MANUAL REVIEW"


def test_top_bottom_layer_assignment_is_not_interchangeable(tmp_path):
    top = _gerber(tmp_path, "top.GTO", _rectangle())
    bottom = _gerber(tmp_path, "bottom.GBO", _rectangle())
    assert derive_gerber_dimension("U1", 10, 20, "Top", [bottom]).length_mm is None
    assert derive_gerber_dimension("U1", 10, 20, "Bottom", [top]).length_mm is None
    assert derive_gerber_dimension("U1", 10, 20, "Top", [top]).length_mm is not None


def test_arc_only_is_not_assumed_rectangular_body(tmp_path):
    doc = _gerber(tmp_path, "arc.GTO",
                  "G75*X80000Y200000D02*G03X120000Y200000I20000J0D01*")
    result = derive_gerber_dimension("U1", 10, 20, "Top", [doc])
    assert result.length_mm is None
    assert result.accepted is False


def test_gui_review_to_excel_export_requires_acceptance(tmp_path):
    from PySide6.QtWidgets import QApplication
    from ui.main_window import MainWindow

    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    part = UniquePart("STEP18-1", ["U1"], "U1")
    cad = CadRecord(ref="U1", x=10, y=20, rotation=0, layer="Top")
    dim = derive_gerber_dimension("U1", 10, 20, "Top",
                                  [_gerber(tmp_path, "body.GTO", _rectangle())])
    assert dim.length_mm is not None and not dim.accepted
    window.state.unique_parts = [part]
    window.state.dimension_results = {part.mpn: dim}
    try:
        window._populate()
        before = tmp_path / "before.xlsx"
        export_excel(before, [part], {"U1": cad}, {part.mpn: dim},
                     {part.mpn: build_shape_model(part, dimension=dim)})
        ws = load_workbook(before, data_only=True)["Shape Dimensions"]
        row = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
        assert row["Body Length (mm)"] is None
        assert row["User Accepted"] == "NO"

        window.table.selectRow(0)
        window._accept_dimension()
        assert dim.accepted and dim.status == "USER ACCEPTED"
        after = tmp_path / "after.xlsx"
        export_excel(after, [part], {"U1": cad}, {part.mpn: dim},
                     {part.mpn: build_shape_model(part, dimension=dim)})
        ws = load_workbook(after, data_only=True)["Shape Dimensions"]
        row = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
        assert row["Body Length (mm)"] == dim.length_mm
        assert row["User Accepted"] == "YES"
    finally:
        window.close()
